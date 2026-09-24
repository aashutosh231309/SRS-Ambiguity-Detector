"""Secure document upload + extraction (API_CONTRACT §4.4, Stage 08) + the
Stage 19 remainder: newest-first list, purge-by-id, and signed-URL downloads.

Covers the full pipeline — validation (filename/extension/MIME/magic/size),
bounded extraction (pdf/docx/txt), upload→analyze integration through the
SHARED Stage 07 engine (equivalence with pasted text), ownership/IDOR, rate
limits, CSRF, timeouts, storage-hygiene (no rows, objects, or temp files
left behind by failures), list pagination + isolation, purge semantics
(analyses survive, pointer degrades), and the download bearer lifecycle
(mint guards, token binding/expiry, byte-identity, attachment headers,
missing/corrupt-object 500s). Self-contained suite (mirrors
test_analysis_crud fixtures; file builders live here — no binary fixtures
in the repo).
"""

import io
import os
import tempfile
import uuid
import zipfile
from collections.abc import Generator
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from httpx import Response
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine

from app.core.config import get_settings
from app.core.database import normalize_url
from app.core.rate_limit import reset_rate_limiter
from app.core.security import create_access_token, create_document_download_token
from app.documents.extraction import (
    DocxExtractor,
    PdfExtractor,
    TxtExtractor,
    extract_text,
)
from app.documents.validation import sanitize_filename, validate_staged_file
from app.exceptions import (
    ExtractedTextTooLargeError,
    ExtractionFailedError,
    FileTooLargeError,
    InvalidFilenameError,
    NoExtractableTextError,
    UnsupportedFileTypeError,
)
from app.main import create_app
from app.models import Analysis, Document, Issue, Requirement, User
from app.storage import LocalStorageBackend
from tests.conftest import TEST_DATABASE_URL, db_test_session, run

# NOTE: credential literals never sit in `password`-named bindings (S106); the
# _pw() helper builds them so every payload value is a call result, not a literal.


@pytest.fixture(autouse=True)
def _fresh_rate_limiter() -> Generator[None, None, None]:
    reset_rate_limiter()
    yield
    reset_rate_limiter()


@pytest.fixture(autouse=True)
def _clean_db(migrated_db: str) -> Generator[None, None, None]:
    _ = migrated_db

    async def _truncate() -> None:
        async with db_test_session():
            pass  # teardown truncates every table (see conftest)

    run(_truncate())
    yield
    run(_truncate())  # app-engine writes bypass db_test_session teardown


@pytest.fixture()
def doc_app(migrated_db: str, tmp_path: Path) -> FastAPI:
    _ = migrated_db
    os.environ["JWT_SECRET"] = "stage08-test-jwt-secret-32-bytes-minimum"  # noqa: S105
    os.environ["STORAGE_LOCAL_DIR"] = str(tmp_path / "uploads")
    get_settings.cache_clear()
    return create_app()


@pytest.fixture()
def doc_client(doc_app: FastAPI) -> Generator[TestClient, None, None]:
    with TestClient(doc_app) as client:
        yield client


@pytest.fixture()
def _retuned_env(monkeypatch: pytest.MonkeyPatch) -> Generator[Any, None, None]:
    """Retune Settings env per-test (request-time resolution picks it up)."""

    def _apply(**pairs: object) -> None:
        for key, value in pairs.items():
            monkeypatch.setenv(key, str(value))
        get_settings.cache_clear()

    yield _apply
    get_settings.cache_clear()


def _email(tag: str) -> str:
    return f"stage08-{tag}-{uuid.uuid4().hex[:8]}@example.com"


def _pw(tag: str = "staple") -> str:
    return f"correct-horse-{tag}-battery-99"


def _register(client: TestClient, email: str) -> Response:
    return client.post(
        "/api/v1/auth/register",
        json={"name": "Stage Eight", "email": email, "password": _pw()},
    )


def _login(client: TestClient, email: str) -> Response:
    return client.post("/api/v1/auth/login", json={"email": email, "password": _pw()})


async def _verify_user(email: str) -> None:
    engine = create_async_engine(normalize_url(TEST_DATABASE_URL))
    try:
        async with AsyncSession(engine) as session:
            user = (await session.execute(select(User).where(User.email == email))).scalar_one()
            user.is_verified = True
            await session.commit()
    finally:
        await engine.dispose()


def _login_verified(client: TestClient, tag: str) -> str:
    email = _email(tag)
    assert _register(client, email).status_code == 201
    run(_verify_user(email))
    assert _login(client, email).status_code == 200
    return email


async def _row_counts() -> tuple[int, int, int, int]:
    """(documents, analyses, requirements, issues) row counts."""
    engine = create_async_engine(normalize_url(TEST_DATABASE_URL))
    try:
        async with AsyncSession(engine) as session:
            counts = []
            for model in (Document, Analysis, Requirement, Issue):
                counts.append(
                    (await session.execute(select(func.count()).select_from(model))).scalar_one()
                )
            return counts[0], counts[1], counts[2], counts[3]
    finally:
        await engine.dispose()


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _make_pdf(pages_text: list[str]) -> bytes:
    """Minimal VALID pdf (catalog/pages/fonts + xref) — parseable by pypdf."""
    n = len(pages_text)
    page_ids = list(range(3, 3 + n))
    content_ids = list(range(3 + n, 3 + 2 * n))
    font_id = 3 + 2 * n
    objs: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{p} 0 R' for p in page_ids)}] /Count {n} >>".encode(),
    ]
    for _page, content in zip(page_ids, content_ids, strict=True):
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
    for text in pages_text:
        stream = f"BT /F1 12 Tf 72 720 Td ({_pdf_escape(text)}) Tj ET".encode()
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    )
    return bytes(out)


def _make_docx(paragraphs: list[str], tables: list[list[list[str]]] | None = None) -> bytes:
    from docx import Document

    document = Document()
    for para in paragraphs:
        document.add_paragraph(para)
    for table in tables or []:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, cells in zip(grid.rows, table, strict=True):
            for cell, text in zip(row.cells, cells, strict=True):
                cell.text = text
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _upload(
    client: TestClient,
    content: bytes,
    filename: str,
    content_type: str = "application/octet-stream",
    title: str | None = None,
) -> Response:
    data = {} if title is None else {"title": title}
    return client.post(
        "/api/v1/documents/upload",
        files={"files": (filename, content, content_type)},
        data=data,
    )


def _code(response: Response) -> str:
    return str(response.json()["error"]["code"])


def _staged_temp_files() -> list[str]:
    return [name for name in os.listdir(tempfile.gettempdir()) if name.startswith("srs-upload-")]


class TestFilenameSanitization:
    @pytest.mark.parametrize(
        ("raw", "expected"),
        [
            ("requirements.docx", "requirements.docx"),
            ("../../etc/passwd.txt", "passwd.txt"),
            ("..\\..\\win\\evil.txt", "evil.txt"),
            ("/abs/path/spec.pdf", "spec.pdf"),
            ("C:\\Users\\ada\\srs.docx", "srs.docx"),
            ("a\x00b.txt", "ab.txt"),
            ("tab\tname.txt", "tabname.txt"),
            ("  padded.txt  ", "padded.txt"),
            ("trailing...", "trailing"),
            ("é.txt", "é.txt"),  # NFD → NFC (visually identical, canonically one)
            ("double.pdf.exe", "double.pdf.exe"),  # content gate decides, not the name
        ],
    )
    def test_neutralizes_without_rejecting(self, raw: str, expected: str) -> None:
        assert sanitize_filename(raw) == expected

    @pytest.mark.parametrize("raw", [None, "", "   ", "...", "/", "///", 12345])
    def test_rejects_unusable_names(self, raw: object) -> None:
        with pytest.raises(InvalidFilenameError):
            sanitize_filename(raw)

    def test_rejects_overlong_names(self) -> None:
        with pytest.raises(InvalidFilenameError):
            sanitize_filename("a" * 252 + ".txt")  # 256 > 255
        assert sanitize_filename("a" * 251 + ".txt") == "a" * 251 + ".txt"  # 255 ok


class TestUploadValidation:
    def test_txt_happy_path_returns_document_plus_analysis(
        self, doc_client: TestClient, tmp_path: Path
    ) -> None:
        _login_verified(doc_client, "a")
        content = b"FR-001: The system shall allow login.\nFR-002: The system shall log out.\n"
        response = _upload(doc_client, content, "srs.txt", "text/plain", title="Txt SRS")
        assert response.status_code == 201
        body = response.json()
        document, analysis = body["document"], body["analysis"]
        assert document["filename"] == "srs.txt"
        assert document["file_type"] == "txt"
        assert document["mime_type"] == "text/plain"
        assert document["byte_size"] == len(content)
        assert document["extracted_chars"] > 0
        assert document["extraction_status"] == "ok"
        assert "storage_path" not in document
        assert analysis["source_type"] == "document"
        assert analysis["status"] == "analyzed"
        assert analysis["title"] == "Txt SRS"
        assert analysis["requirements_count"] == 2
        assert run(_row_counts()) == (1, 1, 2, analysis["issues_count"])
        stored = list((tmp_path / "uploads").rglob("source"))
        assert len(stored) == 1 and stored[0].read_bytes() == content

    def test_docx_and_pdf_happy_paths(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        docx = _make_docx(["Login Requirements", "FR-001: The system shall allow login."])
        response = _upload(
            doc_client,
            docx,
            "srs.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        assert response.status_code == 201
        assert response.json()["document"]["file_type"] == "docx"
        assert response.json()["analysis"]["title"] == "srs.docx"  # filename fallback

        pdf = _make_pdf(["FR-001: The system shall allow login."])
        response = _upload(doc_client, pdf, "srs.pdf", "application/pdf")
        assert response.status_code == 201
        assert response.json()["document"]["file_type"] == "pdf"

    def test_unsupported_extension_rejected(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        response = _upload(doc_client, b"content", "evil.exe")
        assert response.status_code == 400
        assert _code(response) == "unsupported_file_type"
        assert run(_row_counts()) == (0, 0, 0, 0)

    def test_missing_extension_rejected(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        response = _upload(doc_client, b"FR-001: hi", "README")
        assert response.status_code == 400
        assert _code(response) == "unsupported_file_type"

    def test_conflicting_mime_rejected(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        pdf = _make_pdf(["FR-001: The system shall allow login."])
        response = _upload(doc_client, pdf, "srs.pdf", "image/png")
        assert response.status_code == 400
        assert _code(response) == "unsupported_file_type"

    def test_generic_octet_stream_tolerated_when_content_verifies(
        self, doc_client: TestClient
    ) -> None:
        _login_verified(doc_client, "a")  # curl-style: no meaningful MIME claim
        response = _upload(doc_client, b"FR-001: The system shall allow login.\n", "s.txt")
        assert response.status_code == 201

    def test_invalid_pdf_signature_rejected(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        response = _upload(doc_client, b"FR-001: not a pdf at all", "fake.pdf", "application/pdf")
        assert response.status_code == 400
        assert _code(response) == "unsupported_file_type"

    def test_fake_docx_plain_zip_rejected(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as archive:
            archive.writestr("notes.txt", "not ooxml")
        response = _upload(doc_client, buf.getvalue(), "fake.docx")
        assert response.status_code == 400
        assert _code(response) == "unsupported_file_type"

    def test_fake_docx_non_zip_rejected(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        response = _upload(doc_client, b"PK but not a zip" + b"x" * 100, "fake.docx")
        assert response.status_code == 400
        assert _code(response) == "unsupported_file_type"

    def test_binary_renamed_to_txt_rejected(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        response = _upload(doc_client, b"\x7fELF\x00binary", "prog.txt", "text/plain")
        assert response.status_code == 400
        assert _code(response) == "unsupported_file_type"

    def test_empty_file_rejected(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        response = _upload(doc_client, b"", "empty.txt", "text/plain")
        assert response.status_code == 400
        assert _code(response) == "empty_file"

    def test_too_many_files_rejected(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        response = doc_client.post(
            "/api/v1/documents/upload",
            files=[
                ("files", ("a.txt", b"FR-001: one", "text/plain")),
                ("files", ("b.txt", b"FR-002: two", "text/plain")),
            ],
        )
        assert response.status_code == 400
        assert _code(response) == "too_many_files"
        assert response.json()["error"]["details"] == {"max_files": 1}
        assert run(_row_counts()) == (0, 0, 0, 0)

    def test_overlong_title_rejected(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        response = _upload(doc_client, b"FR-001: hi\n", "s.txt", title="t" * 201)
        assert response.status_code == 400
        assert _code(response) == "validation_error"

    def test_traversal_filename_neutralized_not_rejected(
        self, doc_client: TestClient, tmp_path: Path
    ) -> None:
        _login_verified(doc_client, "a")
        response = _upload(doc_client, b"FR-001: The system shall allow login.\n", "../x.txt")
        assert response.status_code == 201
        assert response.json()["document"]["filename"] == "x.txt"
        # Only the server-generated key exists under storage — no traversal wrote out.
        assert [p.name for p in (tmp_path / "uploads").rglob("*") if p.is_file()] == ["source"]

    def test_oversized_bytes_rejected_at_default_limit(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        response = _upload(doc_client, b"a" * (10 * 1024 * 1024 + 1), "big.txt")
        assert response.status_code == 400
        assert _code(response) == "file_too_large"
        assert response.json()["error"]["details"]["reason"] == "byte_size"
        assert run(_row_counts()) == (0, 0, 0, 0)

    def test_byte_limit_is_configurable(self, doc_client: TestClient, _retuned_env: Any) -> None:
        _login_verified(doc_client, "a")
        _retuned_env(MAX_UPLOAD_SIZE_BYTES=16)
        response = _upload(doc_client, b"FR-001: over sixteen bytes", "s.txt")
        assert response.status_code == 400
        assert _code(response) == "file_too_large"

    def test_extracted_text_limit_is_configurable(
        self, doc_client: TestClient, _retuned_env: Any
    ) -> None:
        _login_verified(doc_client, "a")
        _retuned_env(MAX_EXTRACTED_TEXT_CHARS=10)
        response = _upload(doc_client, b"FR-001: way over ten chars", "s.txt")
        assert response.status_code == 400
        assert _code(response) == "extracted_text_too_large"
        assert run(_row_counts()) == (0, 0, 0, 0)

    def test_malformed_pdf_is_422_without_internals(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        response = _upload(
            doc_client, b"%PDF-1.4\ntruncated garbage ((((", "bad.pdf", "application/pdf"
        )
        assert response.status_code == 422
        assert _code(response) == "extraction_failed"
        assert "Traceback" not in response.text and "pypdf" not in response.text
        assert run(_row_counts()) == (0, 0, 0, 0)

    def test_corrupt_docx_xml_is_422(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        good = _make_docx(["FR-001: The system shall allow login."])
        src, dst = io.BytesIO(good), io.BytesIO()
        with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
            for item in zin.infolist():
                data = b"<broken xml" if item.filename == "word/document.xml" else zin.read(item)
                zout.writestr(item, data)
        response = _upload(doc_client, dst.getvalue(), "bad.docx")
        assert response.status_code == 422
        assert _code(response) == "extraction_failed"

    def test_image_only_pdf_reports_no_text(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        response = _upload(doc_client, _make_pdf(["   "]), "scan.pdf", "application/pdf")
        assert response.status_code == 422
        assert _code(response) == "no_extractable_text"

    def test_failures_leave_no_temp_files(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        before = _staged_temp_files()
        _upload(doc_client, b"not a pdf", "fake.pdf", "application/pdf")
        _upload(doc_client, b"%PDF-1.4\njunk", "bad.pdf", "application/pdf")
        _upload(doc_client, b"FR-001: The system shall allow login.\n", "ok.txt")
        assert _staged_temp_files() == before


class TestExtractionUnits:
    """Extractor-level: structure preservation, encodings, caps (no HTTP/DB)."""

    def _staged(self, tmp_path: Path, content: bytes) -> Path:
        path = tmp_path / "staged.bin"
        path.write_bytes(content)
        return path

    def test_pdf_joins_pages_and_reports_count(self, tmp_path: Path) -> None:
        result = PdfExtractor().extract(
            self._staged(tmp_path, _make_pdf(["page one text", "page two text"])),
            max_chars=200_000,
        )
        assert result.text == "page one text\n\npage two text"
        assert result.pages == 2

    def test_pdf_page_cap(self, tmp_path: Path) -> None:
        with pytest.raises(FileTooLargeError) as exc:
            PdfExtractor().extract(
                self._staged(tmp_path, _make_pdf(["p"] * 501)), max_chars=200_000
            )
        assert exc.value.details == {"reason": "pdf_pages", "limit": 500}

    def test_pdf_char_cap_trips_mid_accumulation(self, tmp_path: Path) -> None:
        with pytest.raises(ExtractedTextTooLargeError):
            PdfExtractor().extract(self._staged(tmp_path, _make_pdf(["aaaa", "bbbb"])), max_chars=6)

    def test_pdf_malformed_bytes_raise_extraction_failed(self, tmp_path: Path) -> None:
        with pytest.raises(ExtractionFailedError):
            PdfExtractor().extract(
                self._staged(tmp_path, b"%PDF-1.4\n\x00\x01garbage"), max_chars=200_000
            )

    def test_docx_preserves_document_order(self, tmp_path: Path) -> None:
        from docx import Document

        document = Document()
        document.add_paragraph("FR-001: first requirement.")
        grid = document.add_table(rows=1, cols=2)
        grid.rows[0].cells[0].text = "FR-002"
        grid.rows[0].cells[1].text = "tabled requirement"
        document.add_paragraph("FR-003: last requirement.")
        buf = io.BytesIO()
        document.save(buf)
        result = DocxExtractor().extract(self._staged(tmp_path, buf.getvalue()), max_chars=200_000)
        assert result.text.split("\n") == [
            "FR-001: first requirement.",
            "FR-002 | tabled requirement",
            "FR-003: last requirement.",
        ]

    def test_docx_headings_are_kept_as_text(self, tmp_path: Path) -> None:
        from docx import Document

        document = Document()
        document.add_heading("Scope", level=1)
        document.add_paragraph("FR-001: The system shall allow login.")
        buf = io.BytesIO()
        document.save(buf)
        result = DocxExtractor().extract(self._staged(tmp_path, buf.getvalue()), max_chars=200_000)
        assert result.text.split("\n") == ["Scope", "FR-001: The system shall allow login."]

    def test_docx_empty_is_no_text(self, tmp_path: Path) -> None:
        from docx import Document

        buf = io.BytesIO()
        Document().save(buf)
        with pytest.raises(NoExtractableTextError):
            DocxExtractor().extract(self._staged(tmp_path, buf.getvalue()), max_chars=200_000)

    def test_docx_char_cap(self, tmp_path: Path) -> None:
        with pytest.raises(ExtractedTextTooLargeError):
            DocxExtractor().extract(self._staged(tmp_path, _make_docx(["a" * 50])), max_chars=10)

    def test_txt_decode_chain_and_line_endings(self, tmp_path: Path) -> None:
        # \x93 IS defined in windows-1252 (left smart quote) — latin-1 bytes
        # carrying smart quotes stop at stage 2, not at latin-1.
        smart = TxtExtractor().extract(
            self._staged(tmp_path, "café \x93quoted\x94".encode("latin-1")),
            max_chars=200_000,
        )
        assert smart.text == "café \u201cquoted\u201d"
        assert smart.encoding == "windows-1252"
        # \x81 is UNDEFINED in windows-1252 (one of five such bytes: 81 8D 8F
        # 90 9D) — only the latin-1 final fallback can decode it.
        fallback = TxtExtractor().extract(
            self._staged(tmp_path, b"line one\r\nline two\rna\x81ve"), max_chars=200_000
        )
        assert fallback.text == "line one\nline two\nna\x81ve"
        assert fallback.encoding == "latin-1"

    def test_txt_bom_stripped(self, tmp_path: Path) -> None:
        result = TxtExtractor().extract(
            self._staged(tmp_path, "FR-001: hi".encode("utf-8-sig")), max_chars=200_000
        )
        assert result.text == "FR-001: hi"
        assert result.encoding == "utf-8-sig"

    def test_txt_whitespace_only_is_no_text(self, tmp_path: Path) -> None:
        with pytest.raises(NoExtractableTextError):
            TxtExtractor().extract(self._staged(tmp_path, b"  \n\t\n"), max_chars=200_000)

    def test_txt_char_cap(self, tmp_path: Path) -> None:
        with pytest.raises(ExtractedTextTooLargeError):
            TxtExtractor().extract(self._staged(tmp_path, b"x" * 100), max_chars=10)

    def test_nul_rejected_at_validation_not_extraction(self, tmp_path: Path) -> None:
        staged = self._staged(tmp_path, b"ab\x00cd")
        with pytest.raises(UnsupportedFileTypeError):
            validate_staged_file(
                staged, filename="b.txt", declared_mime="text/plain", byte_size=5, sha256="x"
            )

    def test_extract_text_dispatches_by_validated_type(self, tmp_path: Path) -> None:
        result = extract_text(
            self._staged(tmp_path, b"FR-001: dispatched\n"), file_type="txt", max_chars=200_000
        )
        assert result.file_type == "txt" and "dispatched" in result.text

    def test_docx_member_cap(self, tmp_path: Path) -> None:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as archive:
            archive.writestr("[Content_Types].xml", "x")
            archive.writestr("word/document.xml", "x")
            for i in range(2001):
                archive.writestr(f"word/media/f{i}.bin", "x")
        staged = self._staged(tmp_path, buf.getvalue())
        with pytest.raises(FileTooLargeError) as exc:
            validate_staged_file(
                staged,
                filename="bomb.docx",
                declared_mime="application/octet-stream",
                byte_size=len(buf.getvalue()),
                sha256="x",
            )
        assert exc.value.details == {"reason": "docx_members", "limit": 2000}


class TestUploadIntegration:
    def test_get_document_returns_owned_metadata(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        created = _upload(doc_client, b"FR-001: The system shall allow login.\n", "s.txt")
        document = created.json()["document"]
        response = doc_client.get(f"/api/v1/documents/{document['id']}")
        assert response.status_code == 200
        assert response.json() == document
        assert "storage_path" not in response.text

    def test_same_content_matches_pasted_text_analysis(self, doc_client: TestClient) -> None:
        """Equivalence guarantee: identical text ⇒ identical findings/scores."""
        _login_verified(doc_client, "a")
        text = (
            "FR-001: The system should respond quickly to user requests.\n"
            "FR-002: The service must be available always.\n"
        )
        pasted = doc_client.post("/api/v1/analysis", json={"title": "Pasted", "text": text}).json()
        uploaded = _upload(doc_client, text.encode(), "same.txt", "text/plain").json()["analysis"]

        def _canon(analysis: dict[str, object]) -> object:
            reqs = analysis["requirements"]
            assert isinstance(reqs, list)
            return (
                analysis["score"],
                analysis["band"],
                analysis["issues_count"],
                analysis["health"],
                [
                    (
                        r["text"],
                        r["score"],
                        r["severity"],
                        [(i["detector_id"], i["severity"], i["phrase"]) for i in r["issues"]],
                    )
                    for r in reqs
                ],
            )

        assert _canon(uploaded) == _canon(pasted)
        assert uploaded["source_type"] == "document"
        assert pasted["source_type"] == "text"

    def test_list_filter_source_type_document(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        doc_id = _upload(doc_client, b"FR-001: The system shall allow login.\n", "s.txt").json()[
            "analysis"
        ]["id"]
        doc_client.post(
            "/api/v1/analysis", json={"title": "T", "text": "FR-009: pasted text here."}
        )
        only_docs = doc_client.get("/api/v1/analysis?source_type=document").json()["items"]
        assert [item["id"] for item in only_docs] == [doc_id]
        assert all(item["source_type"] == "document" for item in only_docs)

    def test_delete_analysis_purges_orphaned_document(
        self, doc_client: TestClient, tmp_path: Path
    ) -> None:
        _login_verified(doc_client, "a")
        body = _upload(doc_client, b"FR-001: The system shall allow login.\n", "s.txt").json()
        analysis_id, document_id = body["analysis"]["id"], body["document"]["id"]
        assert doc_client.delete(f"/api/v1/analysis/{analysis_id}").status_code == 204
        assert doc_client.get(f"/api/v1/documents/{document_id}").status_code == 404
        assert run(_row_counts()) == (0, 0, 0, 0)
        assert list((tmp_path / "uploads").rglob("*")) == []

    def test_post_analysis_document_id_still_rejected(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload(
            doc_client, b"FR-001: The system shall allow login.\n", "s.txt"
        ).json()["document"]["id"]
        response = doc_client.post(
            "/api/v1/analysis",
            json={"title": "T", "text": "FR-001: x", "document_id": document_id},
        )
        assert response.status_code == 400
        assert _code(response) == "document_analysis_unavailable"


class TestUploadSecurity:
    def test_unauthenticated_upload_rejected(self, doc_client: TestClient) -> None:
        doc_client.cookies.clear()
        response = _upload(doc_client, b"FR-001: hi\n", "s.txt")
        assert response.status_code == 401
        assert _code(response) == "unauthenticated"

    def test_unverified_user_rejected(self, doc_client: TestClient) -> None:
        email = _email("unverified")
        assert _register(doc_client, email).status_code == 201
        assert _login(doc_client, email).status_code == 200  # login ok, not verified
        response = _upload(doc_client, b"FR-001: hi\n", "s.txt")
        assert response.status_code == 403
        assert _code(response) == "email_unverified"

    def test_foreign_document_is_404(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload(
            doc_client, b"FR-001: The system shall allow login.\n", "a.txt"
        ).json()["document"]["id"]
        doc_client.cookies.clear()
        _login_verified(doc_client, "b")  # re-login overwrites the jar (single client)
        assert doc_client.get(f"/api/v1/documents/{document_id}").status_code == 404
        missing = doc_client.get(f"/api/v1/documents/{document_id}").json()
        assert missing["error"]["code"] == "document_not_found"
        ghost = doc_client.get("/api/v1/documents/00000000-0000-0000-0000-000000000000")
        assert ghost.status_code == 404 and ghost.json() == missing  # identical: no oracle

    def test_bad_origin_rejected(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        response = doc_client.post(
            "/api/v1/documents/upload",
            files={"files": ("s.txt", b"FR-001: hi\n", "text/plain")},
            headers={"Origin": "https://evil.example"},
        )
        assert response.status_code == 403
        assert _code(response) == "forbidden"

    def test_upload_rate_limit(self, doc_client: TestClient, _retuned_env: Any) -> None:
        _login_verified(doc_client, "a")
        _retuned_env(RATE_LIMIT_UPLOADS_PER_MINUTE=2)
        assert _upload(doc_client, b"FR-001: one\n", "a.txt").status_code == 201
        assert _upload(doc_client, b"FR-002: two\n", "b.txt").status_code == 201
        limited = _upload(doc_client, b"FR-003: three\n", "c.txt")
        assert limited.status_code == 429
        assert _code(limited) == "rate_limited"
        assert "Retry-After" in limited.headers

    def test_json_body_to_upload_is_validation_error(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        response = doc_client.post("/api/v1/documents/upload", json={"files": []})
        assert response.status_code == 400
        assert _code(response) == "validation_error"

    def test_processing_timeout_is_503_with_cleanup(
        self, doc_client: TestClient, _retuned_env: Any, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        import app.services.documents as documents_service

        _login_verified(doc_client, "a")
        _retuned_env(DOCUMENT_PROCESSING_TIMEOUT_SECONDS=1)

        def _slow(*args: object, **kwargs: object) -> object:
            import time

            time.sleep(3)
            raise AssertionError("must be abandoned before returning")

        monkeypatch.setattr(documents_service, "_validate_and_extract", _slow)
        before = _staged_temp_files()
        response = _upload(doc_client, b"FR-001: slow parse\n", "slow.txt")
        assert response.status_code == 503
        assert _code(response) == "document_processing_timeout"
        assert run(_row_counts()) == (0, 0, 0, 0)
        assert _staged_temp_files() == before


class TestLocalStorage:
    def test_roundtrip_and_prune(self, tmp_path: Path) -> None:
        backend = LocalStorageBackend(tmp_path / "store")
        src = tmp_path / "src.bin"
        src.write_bytes(b"bytes")
        backend.store_file("documents/u/d/source", src)
        assert not src.exists()  # moved, not copied
        assert (tmp_path / "store" / "documents" / "u" / "d" / "source").read_bytes() == b"bytes"
        backend.delete("documents/u/d/source")
        assert not (tmp_path / "store" / "documents").exists()  # pruned when empty

    def test_delete_missing_is_noop(self, tmp_path: Path) -> None:
        LocalStorageBackend(tmp_path).delete("documents/none/source")

    @pytest.mark.parametrize(
        "key", ["../escape", "/abs/path", "a/../../escape", "..", "a/..\\b", ""]
    )
    def test_traversal_keys_refused(self, tmp_path: Path, key: str) -> None:
        backend = LocalStorageBackend(tmp_path / "store")
        src = tmp_path / "src.bin"
        src.write_bytes(b"x")
        with pytest.raises(ValueError):
            backend.store_file(key, src)

    def test_symlink_escape_refused(self, tmp_path: Path) -> None:
        backend = LocalStorageBackend(tmp_path / "store")
        outside = tmp_path / "outside"
        outside.mkdir()
        link = tmp_path / "store" / "link"
        link.parent.mkdir(parents=True, exist_ok=True)
        link.symlink_to(outside, target_is_directory=True)
        src = tmp_path / "src.bin"
        src.write_bytes(b"x")
        with pytest.raises(ValueError):
            backend.store_file("link/evil", src)

    def test_read_bytes_roundtrip(self, tmp_path: Path) -> None:
        backend = LocalStorageBackend(tmp_path / "store")
        src = tmp_path / "src.bin"
        src.write_bytes(b"\x00\x01binary-bytes")
        backend.store_file("documents/u/d/source", src)
        assert backend.read_bytes("documents/u/d/source") == b"\x00\x01binary-bytes"

    def test_read_bytes_missing_raises(self, tmp_path: Path) -> None:
        with pytest.raises(FileNotFoundError):
            LocalStorageBackend(tmp_path / "store").read_bytes("documents/nope/source")

    def test_read_bytes_refuses_escape(self, tmp_path: Path) -> None:
        with pytest.raises(ValueError):
            LocalStorageBackend(tmp_path / "store").read_bytes("../escape")


class TestDocumentList:
    def test_empty_list_returns_empty_page(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        response = doc_client.get("/api/v1/documents")
        assert response.status_code == 200
        assert response.json() == {"items": [], "page": 1, "page_size": 20, "total": 0}

    def test_list_newest_first_with_pagination(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        ids = [
            _upload(doc_client, b"FR-001: first upload here.\n", "a.txt").json()["document"]["id"],
            _upload(doc_client, b"FR-002: second upload here.\n", "b.txt").json()["document"]["id"],
            _upload(doc_client, b"FR-003: third upload here.\n", "c.txt").json()["document"]["id"],
        ]
        first = doc_client.get("/api/v1/documents?page=1&page_size=2").json()
        assert first["total"] == 3 and first["page"] == 1 and first["page_size"] == 2
        assert [item["id"] for item in first["items"]] == [ids[2], ids[1]]  # newest first
        second = doc_client.get("/api/v1/documents?page=2&page_size=2").json()
        assert [item["id"] for item in second["items"]] == [ids[0]]
        assert second["total"] == 3
        assert "storage_path" not in doc_client.get("/api/v1/documents").text

    def test_list_is_owner_scoped(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        _upload(doc_client, b"FR-001: owner A file here.\n", "a.txt")
        doc_client.cookies.clear()
        _login_verified(doc_client, "b")
        body = doc_client.get("/api/v1/documents").json()
        assert body == {"items": [], "page": 1, "page_size": 20, "total": 0}

    def test_list_requires_verified(self, doc_client: TestClient) -> None:
        assert doc_client.get("/api/v1/documents").status_code == 401
        email = _email("unverified")
        assert _register(doc_client, email).status_code == 201
        assert _login(doc_client, email).status_code == 200
        response = doc_client.get("/api/v1/documents")
        assert response.status_code == 403
        assert _code(response) == "email_unverified"


class TestDocumentDelete:
    def test_delete_purges_row_and_object(self, doc_client: TestClient, tmp_path: Path) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload(
            doc_client, b"FR-001: The system shall allow login.\n", "s.txt"
        ).json()["document"]["id"]
        assert list((tmp_path / "uploads").rglob("*")) != []  # object landed
        assert doc_client.delete(f"/api/v1/documents/{document_id}").status_code == 204
        assert doc_client.get(f"/api/v1/documents/{document_id}").status_code == 404
        assert list((tmp_path / "uploads").rglob("*")) == []  # object + dirs pruned
        docs, analyses, _, _ = run(_row_counts())
        assert (docs, analyses) == (0, 1)  # the analysis SURVIVES the purge

    def test_second_delete_is_404(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload(doc_client, b"FR-001: purge me twice.\n", "s.txt").json()["document"][
            "id"
        ]
        assert doc_client.delete(f"/api/v1/documents/{document_id}").status_code == 204
        missing = doc_client.delete(f"/api/v1/documents/{document_id}")
        assert missing.status_code == 404
        assert _code(missing) == "document_not_found"

    def test_delete_foreign_id_is_identical_404(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload(
            doc_client, b"FR-001: The system shall allow login.\n", "a.txt"
        ).json()["document"]["id"]
        doc_client.cookies.clear()
        _login_verified(doc_client, "b")
        foreign = doc_client.delete(f"/api/v1/documents/{document_id}")
        ghost = doc_client.delete("/api/v1/documents/00000000-0000-0000-0000-000000000000")
        assert foreign.status_code == 404 and ghost.json() == foreign.json()  # no oracle
        assert _code(foreign) == "document_not_found"
        assert run(_row_counts())[0] == 1  # B's attempt changed nothing

    def test_delete_requires_verified(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload(doc_client, b"FR-001: guarded purge.\n", "s.txt").json()["document"][
            "id"
        ]
        doc_client.cookies.clear()
        assert doc_client.delete(f"/api/v1/documents/{document_id}").status_code == 401
        email = _email("unverified")
        assert _register(doc_client, email).status_code == 201
        assert _login(doc_client, email).status_code == 200
        response = doc_client.delete(f"/api/v1/documents/{document_id}")
        assert response.status_code == 403
        assert _code(response) == "email_unverified"

    def test_delete_rejects_bad_origin(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload(doc_client, b"FR-001: csrf purge guard.\n", "s.txt").json()[
            "document"
        ]["id"]
        response = doc_client.delete(
            f"/api/v1/documents/{document_id}", headers={"Origin": "https://evil.example"}
        )
        assert response.status_code == 403
        assert _code(response) == "forbidden"

    def test_linked_analysis_survives_with_null_pointer(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        body = _upload(doc_client, b"FR-001: The system shall allow login.\n", "s.txt").json()
        analysis_id, document_id = body["analysis"]["id"], body["document"]["id"]
        assert body["analysis"]["document"] is not None  # pointer set while linked
        assert doc_client.delete(f"/api/v1/documents/{document_id}").status_code == 204
        detail = doc_client.get(f"/api/v1/analysis/{analysis_id}").json()
        assert detail["document"] is None  # degraded, not broken
        assert detail["score"] == body["analysis"]["score"]  # results intact
        assert detail["requirements"]  # graph intact
        history = doc_client.get("/api/v1/analysis").json()["items"]
        assert history[0]["document"] is None


def _mint(doc_client: TestClient, document_id: str) -> Response:
    return doc_client.post(f"/api/v1/documents/{document_id}/download-url")


def _upload_txt(
    doc_client: TestClient, text: str = "FR-001: The system shall allow login.\n"
) -> str:
    return _upload(doc_client, text.encode(), "s.txt", "text/plain").json()["document"]["id"]


class TestDownloadUrl:
    def test_mint_returns_relative_url_and_expiry(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload_txt(doc_client)
        response = _mint(doc_client, document_id)
        assert response.status_code == 200, response.text
        body = response.json()
        assert set(body) == {"download_url", "expires_at"}
        assert body["download_url"].startswith(f"/api/v1/documents/{document_id}/download?token=")
        assert " " not in body["download_url"]  # single query-safe token
        expires_at = datetime.fromisoformat(body["expires_at"])
        assert expires_at.tzinfo is not None
        now = datetime.now(UTC)
        assert timedelta(minutes=14) < expires_at - now <= timedelta(minutes=15)

    def test_mint_foreign_or_missing_is_identical_404(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload_txt(doc_client)
        doc_client.cookies.clear()
        _login_verified(doc_client, "b")
        foreign = _mint(doc_client, document_id)
        ghost = _mint(doc_client, "00000000-0000-0000-0000-000000000000")
        assert foreign.status_code == 404 and ghost.json() == foreign.json()  # no oracle
        assert _code(foreign) == "document_not_found"

    def test_mint_requires_verified(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload_txt(doc_client)
        doc_client.cookies.clear()
        assert _mint(doc_client, document_id).status_code == 401
        email = _email("unverified")
        assert _register(doc_client, email).status_code == 201
        assert _login(doc_client, email).status_code == 200
        response = _mint(doc_client, document_id)
        assert response.status_code == 403
        assert _code(response) == "email_unverified"

    def test_mint_rejects_bad_origin(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload_txt(doc_client)
        response = doc_client.post(
            f"/api/v1/documents/{document_id}/download-url",
            headers={"Origin": "https://evil.example"},
        )
        assert response.status_code == 403
        assert _code(response) == "forbidden"

    def test_mint_rate_limit(self, doc_client: TestClient, _retuned_env: Any) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload_txt(doc_client)
        _retuned_env(RATE_LIMIT_DOCUMENT_DOWNLOAD_PER_MINUTE=2)
        assert _mint(doc_client, document_id).status_code == 200
        assert _mint(doc_client, document_id).status_code == 200
        limited = _mint(doc_client, document_id)
        assert limited.status_code == 429
        assert _code(limited) == "rate_limited"
        assert "Retry-After" in limited.headers

    def test_download_streams_original_bytes_without_session(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        content = _make_pdf(["FR-001: The system shall allow login quickly."])
        document_id = _upload(doc_client, content, "spec.pdf").json()["document"]["id"]
        url = _mint(doc_client, document_id).json()["download_url"]
        doc_client.cookies.clear()  # the token IS the credential — no session needed
        response = doc_client.get(url)
        assert response.status_code == 200
        assert response.content == content  # byte-identical, not re-rendered
        assert response.headers["content-type"] == "application/pdf"  # detected MIME
        assert response.headers["content-disposition"] == (
            "attachment; filename=\"spec.pdf\"; filename*=UTF-8''spec.pdf"
        )
        assert response.headers["content-length"] == str(len(content))
        assert response.headers["x-content-type-options"] == "nosniff"

    def test_download_names_unicode_filenames_safely(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload(doc_client, b"FR-001: hi there friend.\n", "srs-ünïcödé.txt").json()[
            "document"
        ]["id"]
        url = _mint(doc_client, document_id).json()["download_url"]
        response = doc_client.get(url)
        assert response.status_code == 200
        assert response.headers["content-disposition"] == (
            'attachment; filename="srs-?n?c?d?.txt"; '
            "filename*=UTF-8''srs-%C3%BCn%C3%AFc%C3%B6d%C3%A9.txt"
        )

    def test_download_rejects_expired_token(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload_txt(doc_client)
        stale = create_document_download_token(uuid.uuid4(), uuid.UUID(document_id), -1)
        response = doc_client.get(f"/api/v1/documents/{document_id}/download?token={stale}")
        assert response.status_code == 400
        assert _code(response) == "invalid_token"

    def test_download_rejects_forged_token(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload_txt(doc_client)
        url = _mint(doc_client, document_id).json()["download_url"]
        token = url.rsplit("token=", 1)[1]
        forged = token[:-1] + ("a" if token[-1] != "a" else "b")
        response = doc_client.get(f"/api/v1/documents/{document_id}/download?token={forged}")
        assert response.status_code == 400
        assert _code(response) == "invalid_token"

    def test_download_rejects_cross_document_token(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        doc_a = _upload_txt(doc_client, "FR-001: first document here.\n")
        doc_b = _upload_txt(doc_client, "FR-002: second document here.\n")
        url_a = _mint(doc_client, doc_a).json()["download_url"]
        token_a = url_a.rsplit("token=", 1)[1]
        response = doc_client.get(f"/api/v1/documents/{doc_b}/download?token={token_a}")
        assert response.status_code == 400
        assert _code(response) == "invalid_token"

    def test_download_rejects_session_jwt(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload_txt(doc_client)
        session_jwt = create_access_token(uuid.uuid4())  # valid signature, wrong type
        response = doc_client.get(f"/api/v1/documents/{document_id}/download?token={session_jwt}")
        assert response.status_code == 400
        assert _code(response) == "invalid_token"

    def test_download_after_delete_is_404(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload_txt(doc_client)
        url = _mint(doc_client, document_id).json()["download_url"]
        assert doc_client.delete(f"/api/v1/documents/{document_id}").status_code == 204
        response = doc_client.get(url)
        assert response.status_code == 404
        assert _code(response) == "document_not_found"

    def test_download_missing_object_is_500(self, doc_client: TestClient, tmp_path: Path) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload_txt(doc_client)
        url = _mint(doc_client, document_id).json()["download_url"]
        objects = [p for p in (tmp_path / "uploads").rglob("*") if p.is_file()]
        assert len(objects) == 1
        objects[0].unlink()  # row lives, bytes lost: OUR inconsistency
        response = doc_client.get(url)
        assert response.status_code == 500
        assert _code(response) == "internal_error"
        assert str(tmp_path) not in response.text  # no paths leak

    def test_download_corrupt_object_is_500(self, doc_client: TestClient, tmp_path: Path) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload_txt(doc_client)
        url = _mint(doc_client, document_id).json()["download_url"]
        objects = [p for p in (tmp_path / "uploads").rglob("*") if p.is_file()]
        assert len(objects) == 1
        objects[0].write_bytes(b"tampered-not-the-upload")
        response = doc_client.get(url)
        assert response.status_code == 500
        assert _code(response) == "internal_error"

    def test_download_requires_token(self, doc_client: TestClient) -> None:
        _login_verified(doc_client, "a")
        document_id = _upload_txt(doc_client)
        assert doc_client.get(f"/api/v1/documents/{document_id}/download").status_code == 400
        empty = doc_client.get(f"/api/v1/documents/{document_id}/download?token=")
        assert empty.status_code == 400
        assert _code(empty) == "invalid_token"

    def test_full_lifecycle_roundtrip(self, doc_client: TestClient) -> None:
        # Upload → list shows it → mint → download bytes → purge → gone
        # everywhere (list, metadata, analysis pointer), token dead after.
        _login_verified(doc_client, "a")
        content = b"FR-001: The system shall allow login.\n"
        uploaded = _upload(doc_client, content, "roundtrip.txt", "text/plain").json()
        document_id, analysis_id = uploaded["document"]["id"], uploaded["analysis"]["id"]
        assert [item["id"] for item in doc_client.get("/api/v1/documents").json()["items"]] == [
            document_id
        ]
        minted = _mint(doc_client, document_id).json()
        token = minted["download_url"].rsplit("token=", 1)[1]
        assert token not in doc_client.get("/api/v1/documents").text  # bearer stays in the URL
        assert token not in doc_client.get(f"/api/v1/documents/{document_id}").text
        assert doc_client.get(minted["download_url"]).content == content
        assert doc_client.delete(f"/api/v1/documents/{document_id}").status_code == 204
        assert doc_client.get("/api/v1/documents").json()["items"] == []
        assert doc_client.get(minted["download_url"]).status_code == 404
        detail = doc_client.get(f"/api/v1/analysis/{analysis_id}").json()
        assert detail["document"] is None and detail["score"] is not None
